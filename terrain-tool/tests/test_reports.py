"""Report renderer tests: content blocks, PDF (ReportLab) and DOCX (python-docx) output.

No network. Figures are placeholder PNGs except in the engine integration test, which
runs the real terrain engine on a synthetic DEM. Needs pypdf (requirements-dev.txt).
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
from report_helpers import (FIGURE_SIZES, docx_text, load_sample, pdf_text,  # noqa: E402
                            report_settings, squash, write_placeholder_figures)

from docx import Document  # noqa: E402
from docx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from app.reports import DOCX_NAME, PDF_NAME, build_reports  # noqa: E402
from app.reports import content as c  # noqa: E402

CONSULT = "https://www.ecoconsultancy.services/consult/"
H1_ORDER = ["About this report", "At a glance", "Findings", "Site statistics", "Figures",
            "Recommended follow-up studies", "Caveats and limitations",
            "Data sources and attribution", "Disclaimer", "Next step"]
HOSTILE = '<b>&</b> "x"'


def _build(tmp_path: Path, result: dict | None = None, figures: bool = True, **settings_kw):
    out = tmp_path / "job"
    out.mkdir(parents=True, exist_ok=True)
    if figures:
        write_placeholder_figures(out)
    result = result if result is not None else load_sample()
    paths = build_reports(result, out, report_settings(tmp_path, **settings_kw))
    return result, out, paths


@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    """The sample result rendered once; tests below only read the outputs."""
    tmp = tmp_path_factory.mktemp("reports")
    result, out, paths = _build(tmp)
    return {"result": result, "out": out, "paths": paths,
            "pdf": PdfReader(str(paths["pdf"])), "pdf_text": pdf_text(paths["pdf"]),
            "docx": Document(str(paths["docx"]))}


# ---------------------------------------------------------------------------
# build_reports contract
# ---------------------------------------------------------------------------
def test_build_reports_returns_both_files(sample):
    paths, out = sample["paths"], sample["out"]
    assert set(paths) == {"pdf", "docx"}
    assert paths["pdf"] == out / PDF_NAME == out / "report.pdf"
    assert paths["docx"] == out / DOCX_NAME == out / "report.docx"
    pdf_bytes = paths["pdf"].read_bytes()
    assert pdf_bytes.startswith(b"%PDF-") and pdf_bytes.rstrip().endswith(b"%%EOF")
    assert len(pdf_bytes) > 40_000                    # four embedded figures, not an empty shell
    assert zipfile.is_zipfile(paths["docx"])
    assert paths["docx"].stat().st_size > 40_000
    assert not list(out.glob(".*.tmp"))               # atomic write leaves nothing behind


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def test_pdf_pages_metadata_and_furniture(sample):
    reader, text, meta = sample["pdf"], sample["pdf_text"], sample["result"]["meta"]
    n = len(reader.pages)
    assert 7 <= n <= 14
    info = reader.metadata
    assert info.title == "Desktop Terrain Screening — Kuala Langat LSS candidates"
    assert info.author == "EcoEnergy Consultancy"
    assert meta["report_id"] in info.subject
    # running header on every page with "page N of M", footer label on every page
    short = meta["report_id"][:8]
    for i, page in enumerate(reader.pages, 1):
        page_text = squash(page.extract_text())
        assert f"Report {short} · page {i} of {n}" in page_text
        assert "EcoEnergy Consultancy · Desktop Terrain Screening" in page_text
        assert "Screening grade — not a survey" in page_text
    assert "13 October 2026, 16:00 MYT" in text      # 08:00 UTC shown in Asia/Kuala_Lumpur time
    assert "SCREENING GRADE — NOT A SURVEY" in text


def test_pdf_contains_every_section_and_engine_text(sample):
    text, result = sample["pdf_text"], sample["result"]
    for heading in H1_ORDER + ["Sites", "Slope distribution", "Routes", "Points"]:
        assert heading in text, heading
    for flag in result["flags"]:
        assert flag["title"] in text
    for item in result["followups"]:
        assert item["title"] in text
    for fig in result["figures"]:
        assert fig["title"] in text
    for needle in ("Kuala Langat LSS candidates", "Kuala Langat, Selangor", "Large Scale Solar (LSS)",
                   "Lot 4521 (oil palm, north)", "33 kV cable route to PMU", "Proposed PMU tie-in",
                   "60.97 ha", "Class I 94.2%", "Licensed Land Surveyors Act 1958",
                   "Book a free 30-minute scheme-fit consult",
                   "Analysis: EcoEnergy terrain screening engine v1", "Copernicus DEM GLO-30 © DLR",
                   "geoBoundaries (gbOpen)", "not for submission to any authority"):
        assert needle in text, needle


def test_pdf_cta_link_figures_and_outline(sample):
    reader = sample["pdf"]
    uris = [a.get_object()["/A"]["/URI"] for page in reader.pages for a in page.get("/Annots", []) or []
            if a.get_object().get("/Subtype") == "/Link"]
    expected = CONSULT + "?utm_source=terrain_tool&utm_medium=pdf"
    assert expected in uris
    assert set(uris) == {expected}
    assert sum(len(page.images) for page in reader.pages) == 4
    # each figure sits on its own page
    figure_pages = [i for i, page in enumerate(reader.pages) if page.images]
    assert len(figure_pages) == 4 and len(set(figure_pages)) == 4
    titles = []

    def walk(items):
        for item in items:
            if isinstance(item, list):
                walk(item)
            else:
                titles.append(item.title)
    walk(reader.outline)
    assert [t for t in titles if t in H1_ORDER] == H1_ORDER


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------
def test_docx_headings_tables_and_images(sample):
    doc = sample["docx"]
    h1 = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert h1 == H1_ORDER
    assert h2 == ["Sites", "Slope distribution", "Routes", "Points", "Figure 1. Site layout",
                  "Figure 2. Elevation and contours", "Figure 3. Slope classes",
                  "Figure 4. Terrain profiles"]
    assert len(doc.inline_shapes) == 4
    for shape in doc.inline_shapes:                   # 16 cm wide, or fitted to 20 cm height
        assert (shape.width.cm == pytest.approx(16.0, abs=0.01)
                or shape.height.cm == pytest.approx(20.0, abs=0.01))

    def table_starting(*cells):
        for t in doc.tables:
            first = [cell.text for cell in t.rows[0].cells]
            if all(x in first for x in cells):
                return t
        raise AssertionError(f"no table with header {cells}")

    sites = table_starting("Site", "Area (ha)", "Slope (°)")
    assert [cell.text for cell in sites.rows[1].cells][4:7] == ["Mean", "P90", "Max"]
    assert sites.rows[2].cells[0].text == "Lot 4521 (oil palm, north)"
    assert sites.rows[-1].cells[0].text == "All sites"
    fills = set(sites._tbl.xpath("./w:tr[1]//w:shd/@w:fill"))
    assert fills == {"0F2742"}                        # navy header row
    assert sites._tbl.xpath("./w:tr[1]/w:trPr/w:tblHeader")   # header repeats across pages
    table_starting("Site", "Share of site area in each slope band (%)")
    table_starting("Route", "Length (km)")
    table_starting("Point", "Latitude (°)")


def test_docx_text_header_footer_link_and_properties(sample):
    doc, result = sample["docx"], sample["result"]
    text = docx_text(doc)
    for flag in result["flags"]:
        assert flag["title"] in text
    for needle in ("13 October 2026, 16:00 MYT", "SCREENING GRADE — NOT A SURVEY",
                   "Licensed Land Surveyors Act 1958", "Book a free 30-minute scheme-fit consult",
                   "Analysis: EcoEnergy terrain screening engine v1"):
        assert needle in text, needle
    header_xml = doc.sections[0].header._element.xml
    assert "EcoEnergy Consultancy" in header_xml and "Report q3VtYp7c" in header_xml
    assert " PAGE " in header_xml and " NUMPAGES " in header_xml
    assert "Screening grade — not a survey" in doc.sections[0].footer._element.xml
    links = [r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK]
    assert CONSULT + "?utm_source=terrain_tool&utm_medium=docx" in links
    cp = doc.core_properties
    assert cp.author == "EcoEnergy Consultancy"
    assert cp.title == "Desktop Terrain Screening — Kuala Langat LSS candidates"
    assert result["meta"]["report_id"] in cp.subject
    sec = doc.sections[0]
    assert round(sec.page_width.mm) == 210 and round(sec.page_height.mm) == 297


def test_both_formats_render_the_same_blocks(sample):
    content = c.build_content(sample["result"], sample["out"], report_settings(sample["out"]))
    pdf, docx = sample["pdf_text"], docx_text(sample["docx"])
    for block in content.blocks:
        texts = []
        if isinstance(block, (c.Heading, c.Paragraph)):
            texts = [block.text]
        elif isinstance(block, c.Callout):
            texts = [block.title]
        elif isinstance(block, c.KeyValueTable):
            texts = [v for row in block.rows for v in row]
        for t in texts:
            assert t in pdf, t
            assert t in docx, t


# ---------------------------------------------------------------------------
# robustness
# ---------------------------------------------------------------------------
def test_hostile_and_unicode_names_render_literally(tmp_path):
    result = load_sample()
    result["meta"]["site_label"] = HOSTILE + " </para><font size=40>"
    result["sites"][0]["name"] = HOSTILE
    result["sites"][1]["name"] = 'Lot 地块 7 Ōmura <a href="javascript:alert(1)">x</a>'
    result["routes"][0]["name"] = "bell\x07 and null\x00 route"
    result["points"][0]["name"] = "<img src=x onerror=alert(1)>"
    result["flags"][0]["detail"] += f" {HOSTILE}"
    _, _, paths = _build(tmp_path, result)

    text = pdf_text(paths["pdf"])
    assert HOSTILE in text
    assert "<img src=x onerror=alert(1)>" in text
    assert "Lot" in text and "Omura" in text          # accent dropped for the built-in font
    assert "</para><font size=40>" in text
    reader = PdfReader(str(paths["pdf"]))
    uris = {a.get_object()["/A"]["/URI"] for page in reader.pages for a in page.get("/Annots", []) or []}
    assert all(u.startswith(CONSULT) for u in uris)   # no link smuggled in through a name

    doc = Document(str(paths["docx"]))
    dtext = docx_text(doc)
    assert HOSTILE in dtext
    assert 'Lot 地块 7 Ōmura <a href="javascript:alert(1)">x</a>' in dtext   # Word keeps full Unicode
    assert "bell and null route" in dtext             # XML-illegal control characters stripped
    links = [r.target_ref for r in doc.part.rels.values() if r.reltype == RT.HYPERLINK]
    assert all(u.startswith(CONSULT) for u in links)


def test_missing_values_figures_and_sections(tmp_path):
    result = load_sample()
    result["meta"].update(site_label=None, scheme=None, stage=None)
    result["location"].update(state=None, district=None)
    result["dem"]["grid_spacing_m"] = None
    site = result["sites"][0]
    for key in ("elev_min_m", "elev_max_m", "relief_m",
                "slope_mean_deg", "slope_p90_deg", "slope_max_deg"):
        site[key] = None
    site["kpkt_pct"] = {"I": None, "II": None, "III": None, "IV": None}
    result["routes"], result["points"] = [], []
    result["figures"][1]["file"] = "../../etc/passwd"   # must never be read
    result["figures"][2]["file"] = "missing.png"
    out = tmp_path / "job"
    write_placeholder_figures(out, ["1_site_layout.png", "4_profiles.png"])
    (out / "4_profiles.png").write_bytes(b"not a png")  # unreadable image -> placeholder
    _, _, paths = _build(tmp_path, result, figures=False)

    text = pdf_text(paths["pdf"])
    assert "kuala_langat_candidates.kmz" in text       # filename stands in for the missing label
    assert "State and district not identified" in text
    assert "Not stated" in text
    assert text.count("This figure could not be produced.") == 3
    assert "—" in text
    doc = Document(str(paths["docx"]))
    assert len(doc.inline_shapes) == 1
    h2 = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert "Routes" not in h2 and "Points" not in h2
    assert "Sites" in h2


def test_many_sites_and_long_names_flow_across_pages(tmp_path):
    result = load_sample()
    base = result["sites"]
    result["sites"] = [dict(base[i % 2], name=f"Parcel {i:02d} " + "x" * (300 if i == 3 else 5))
                       for i in range(50)]
    result["sites"][7]["name"] = "N" * 500             # no spaces at all
    _, _, paths = _build(tmp_path, result)
    reader = PdfReader(str(paths["pdf"]))
    assert len(reader.pages) > 10
    text = pdf_text(paths["pdf"])
    assert "Parcel 49" in text and "…" in text
    doc = Document(str(paths["docx"]))
    sites = next(t for t in doc.tables if t.rows[0].cells[0].text == "Site")
    assert len(sites.rows) == 2 + 50 + 1


def test_rebuild_overwrites_existing_reports(tmp_path):
    _, out, first = _build(tmp_path)
    first_pdf = first["pdf"].read_bytes()
    result = load_sample()
    result["meta"]["site_label"] = "Second run"
    second = build_reports(result, out, report_settings(tmp_path))
    assert second == first
    assert second["pdf"].read_bytes() != first_pdf
    assert "Second run" in pdf_text(second["pdf"])


# ---------------------------------------------------------------------------
# content helpers
# ---------------------------------------------------------------------------
def test_tracked_url():
    assert c.tracked_url(CONSULT, "pdf") == CONSULT + "?utm_source=terrain_tool&utm_medium=pdf"
    assert (c.tracked_url("https://x.test/c/?ref=a&utm_medium=old#book", "docx")
            == "https://x.test/c/?ref=a&utm_source=terrain_tool&utm_medium=docx#book")


def test_formatting_helpers():
    assert c.format_generated("2026-10-13T08:00:00+00:00") == "13 October 2026, 16:00 MYT"
    assert c.format_generated("2026-12-31T20:30:00Z") == "1 January 2027, 04:30 MYT"
    assert c.format_generated("not a date") == "not a date"
    assert c.format_generated(None) == "—"
    assert c.num(None) == c.num(float("nan")) == c.num(True) == "—"
    assert c.num(1234.567, 2, " ha") == "1,234.57 ha"
    assert c.num_range(None, None) == "—" and c.num_range(1, None) == "1.0–—"
    assert c.bin_label("0-2") == "0–2°" and c.bin_label(">35") == ">35°"
    assert c.name_text("  a \n b  ") == "a b" and c.name_text(None) == "—"
    assert len(c.name_text("y" * 1000)) == c.MAX_NAME_CHARS
    assert c.clean("ok\x00\x1f￾") == "ok"


def test_content_order_and_cta(tmp_path):
    write_placeholder_figures(tmp_path)
    content = c.build_content(load_sample(), tmp_path, report_settings(tmp_path))
    blocks = content.blocks
    assert isinstance(blocks[0], c.Cover)
    assert [b.text for b in blocks if isinstance(b, c.Heading) and b.level == 1] == H1_ORDER
    assert [b.level for b in blocks if isinstance(b, c.Callout)] == \
        ["warning", "warning", "caution", "caution", "info", "info", "info"]
    figures = [b for b in blocks if isinstance(b, c.Figure)]
    assert [f.path.name for f in figures] == list(FIGURE_SIZES)
    assert all(f.size_px and f.size_px[0] > 0 for f in figures)
    cta = blocks[-1]
    assert isinstance(cta, c.Cta) and cta.title == "Book a free 30-minute scheme-fit consult"
    assert cta.url_for("pdf").endswith("utm_medium=pdf")
    assert cta.display_url == "www.ecoconsultancy.services/consult/"
    assert content.info.short_id == "q3VtYp7c"


# ---------------------------------------------------------------------------
# integration with the real terrain engine (synthetic DEM, no network)
# ---------------------------------------------------------------------------
def test_reports_from_real_engine_output(tmp_path):
    import numpy as np
    from engine_helpers import (function_fetcher, kml_doc, line_pm, offset, point_pm, polygon_pm,
                                rect, write)

    from app.terrain import run_screening

    lon, lat = 101.45, 2.86
    body = (polygon_pm("Lot A", rect(lon, lat, 700, 600))
            + polygon_pm("Hill <B>", rect(*offset(lon, lat, 900, 100), 450, 420))
            + line_pm("Cable route", [offset(lon, lat, -400, -300), offset(lon, lat, 1100, 300)])
            + point_pm("Tie-in", offset(lon, lat, -400, -300)))
    kml = write(tmp_path, "site.kml", kml_doc(body.replace("<B>", "&lt;B&gt;")))

    def terrain(LON, LAT, X, Y):
        return 15 + 0.004 * X + 60 * np.exp(-(((X - 700) / 260) ** 2 + (Y / 260) ** 2))

    settings = report_settings(tmp_path)
    meta = {"report_id": "engineTest0001", "generated_at": "2026-10-13T08:00:00+00:00",
            "site_label": None, "scheme": "lss", "stage": None, "input_filename": "site.kml"}
    out = tmp_path / "job"
    result = run_screening(kml, out, settings, meta, dem_fetcher=function_fetcher(terrain))
    paths = build_reports(result, out, settings)

    text = pdf_text(paths["pdf"])
    for needle in ("Lot A", "Hill <B>", "Cable route", "Tie-in", "Report engineTe"):
        assert needle in text, needle
    for flag in result["flags"]:
        assert flag["title"] in text
    assert sum(len(p.images) for p in PdfReader(str(paths["pdf"])).pages) == 4
    assert len(Document(str(paths["docx"])).inline_shapes) == 4
